"""
导出模块

把分析结果导出为**可直接投稿**的三类交付物：

- Word 报告：统计表格按中英文期刊惯例排成**三线表**，插图为 300 dpi
- Excel 结果表：每个章节一个工作表，便于审稿人复核原始数字
- 图片包 ZIP：300 dpi PNG + 带 dpi 元数据的 TIFF（附图目录）

说明：绘图分辨率由各分析模块在 savefig 时指定为 300 dpi（期刊一般要求 ≥300 dpi），
本模块只负责排版与打包，不重新渲染图片。
"""
import io
import math
import re
import zipfile
from datetime import datetime
from typing import Dict, List, Optional

import pandas as pd
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

# 正文可用宽度（A4 纵向、常规页边距约 6.3 英寸）
NARROW_IMAGE_INCHES = 5.5
WIDE_IMAGE_INCHES = 6.2
WIDE_IMAGE_RATIO = 1.7      # 宽高比超过该值视为宽图，用满正文宽度以免被压得太小
THIN_LINE = '6'             # OOXML 的 w:sz 以 1/8 pt 为单位：6 = 0.75 pt
THICK_LINE = '12'           # 12 = 1.5 pt
_SHEET_BAD_CHARS = re.compile(r'[\\/:*?\[\]]')
_FILE_BAD_CHARS = re.compile(r'[\\/:*?"<>|\s]+')


def _safe_sheet_name(title, used_names: set) -> str:
    """Excel 工作表名：去非法字符、限 31 字符并去重。"""
    name = _SHEET_BAD_CHARS.sub('', str(title or '结果')).strip() or '结果'
    name = name[:31]
    base, index = name, 2
    while name in used_names:
        suffix = f'_{index}'
        name = base[:31 - len(suffix)] + suffix
        index += 1
    used_names.add(name)
    return name


def _safe_file_stem(title, index: int) -> str:
    name = _FILE_BAD_CHARS.sub('_', str(title or '图')).strip('_')
    return f'{index:02d}_{name[:50] or "图"}'


def format_cell(value, column: str = '', style=None) -> str:
    """
    表格单元格的数字格式（论文惯例）。

    实际规则见 modules/journal.py：传入列名后，p 值列 / 百分比列 / 计数列会按
    期刊惯例分别处理；不传列名或样式时使用默认规则，行为与旧版本一致。
    """
    from modules.journal import format_value

    return format_value(value, column, style)


def _image_width_inches(data: bytes, style=None) -> float:
    """
    插图嵌入宽度。

    模板指定了宽度就按模板（期刊常按单栏 8.5 cm / 双栏 17 cm 要求）；
    否则按宽高比自适应：宽图用满正文宽度，避免在 Word 里被压得看不清。
    """
    if style is not None and getattr(style, 'image_width_cm', 0):
        return float(style.image_width_cm) / 2.54
    try:
        from PIL import Image

        with Image.open(io.BytesIO(data)) as image:
            width, height = image.size
        if height and width / height >= WIDE_IMAGE_RATIO:
            return WIDE_IMAGE_INCHES
    except Exception:
        pass
    return NARROW_IMAGE_INCHES


def _add_caption(doc, label: str, title: str, style) -> None:
    """
    插入题注，格式为「图 1　标题」。

    期刊惯例：表题在表格**上方**，图题在图片**下方**（由调用方决定插入顺序）。
    """
    text = f'{label}　{title}' if title else label
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(text)
    run.font.size = Pt(float(getattr(style, 'caption_size_pt', 9.0)))
    run.font.bold = bool(getattr(style, 'caption_bold', False))


def _to_tiff(png_bytes: bytes, dpi: int) -> Optional[bytes]:
    """
    把 PNG 转成带 dpi 元数据的 TIFF。

    像素数不变，只是换成投稿常用的容器格式并写入 dpi 信息
    （不少期刊系统按元数据判断分辨率是否达标）。
    """
    try:
        from PIL import Image
    except Exception:
        return None
    try:
        with Image.open(io.BytesIO(png_bytes)) as image:
            converted = image.convert('RGB')
            output = io.BytesIO()
            converted.save(output, format='TIFF', dpi=(dpi, dpi),
                           compression='tiff_lzw')
            return output.getvalue()
    except Exception:
        return None


# OOXML schema 要求子元素按固定顺序出现：tblBorders 必须排在 w:shd / w:tblLayout / w:tblLook 之前，
# tcBorders 必须排在 w:shd / w:vAlign 之前 —— 顺序错了 Word 会整段忽略边框定义，
# 表格就会退化成无框表格。python-docx 没有为这两个元素提供 get_or_add_*，故手动按序插入。
_TBL_BORDERS_SUCCESSORS = ('w:shd', 'w:tblLayout', 'w:tblCellMar', 'w:tblLook',
                           'w:tblCaption', 'w:tblDescription', 'w:tblPrChange')
_TC_BORDERS_SUCCESSORS = ('w:shd', 'w:noWrap', 'w:tcMar', 'w:textDirection',
                          'w:tcFitText', 'w:vAlign', 'w:hideMark')


def _ordered_insert(parent, element, successors) -> None:
    """把 element 插入到第一个 successors 之前，保持 OOXML 的元素顺序。"""
    successor_tags = {qn(tag) for tag in successors}
    for child in parent:
        if child.tag in successor_tags:
            child.addprevious(element)
            return
    parent.append(element)


def _get_or_create_ordered(parent, tag, successors):
    """取子元素；不存在时按 schema 顺序新建。"""
    element = parent.find(qn(tag))
    if element is None:
        element = OxmlElement(tag)
        _ordered_insert(parent, element, successors)
    return element


class ReportExporter:
    """报告导出"""

    @staticmethod
    def export_word(title: str, sections: List[Dict],
                    filename: str = 'analysis_report.docx',
                    table_style=None) -> bytes:
        """
        导出 Word 报告（三线表 + 300 dpi 插图）

        sections:     [{'title': '...', 'content': '文本', 'table': df或None,
                        'image': bytes或None, 'level': 1/2/3}, ...]
        table_style   期刊模板（modules/journal.py 的 TableStyle）；None 表示用默认模板
        """
        from modules.journal import TableStyle
        from modules.localize import caption_prefix, frame_text, translate_title

        table_style = table_style or TableStyle()
        language = str(getattr(table_style, 'language', 'zh') or 'zh')
        table_prefix, figure_prefix = caption_prefix(language)
        doc = Document()

        # 正文字体按模板设置（中文期刊常用宋体，SCI 常用 Times New Roman）
        normal_style = doc.styles['Normal']
        normal_style.font.name = table_style.font
        normal_style.element.rPr.rFonts.set(qn('w:eastAsia'), table_style.font)

        # 标题
        heading = doc.add_heading(title, level=0)
        heading.alignment = WD_ALIGN_PARAGRAPH.CENTER

        # 元信息
        separator = '：' if language == 'zh' else ': '
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(
            f"{frame_text('generated_at', language)}{separator}"
            f"{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(128, 128, 128)

        doc.add_paragraph()

        table_index = 0
        figure_index = 0
        number_captions = bool(getattr(table_style, 'number_captions', True))

        for section in sections:
            level = section.get('level', 1)
            # 英文报告时章节名按术语表翻译（未收录的词保留原文）
            title = translate_title(section.get('title', ''), language)
            doc.add_heading(title, level=level)

            if section.get('content'):
                paragraph = doc.add_paragraph(section['content'])
                paragraph.paragraph_format.first_line_indent = Inches(0.3)

            if section.get('table') is not None:
                table_index += 1
                # 表题在表格上方（中文期刊与 SCI 惯例一致）
                if number_captions:
                    _add_caption(doc, f'{table_prefix} {table_index}', title, table_style)
                ReportExporter._add_df_table(doc, section['table'], table_style)

            if section.get('image'):
                image_stream = io.BytesIO(section['image'])
                doc.add_picture(
                    image_stream,
                    width=Inches(_image_width_inches(section['image'], table_style)))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                figure_index += 1
                # 图题在图片下方
                if number_captions:
                    _add_caption(doc, f'{figure_prefix} {figure_index}', title, table_style)

        # 落款
        doc.add_paragraph()
        paragraph = doc.add_paragraph(frame_text('footer', language))
        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        for run in paragraph.runs:
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(150, 150, 150)
            run.italic = True

        buffer = io.BytesIO()
        doc.save(buffer)
        return buffer.getvalue()

    @staticmethod
    def export_excel(sections: List[Dict]) -> bytes:
        """把各章节的统计表格导出为一个 xlsx，每个章节一个工作表。"""
        buffer = io.BytesIO()
        used_names: set = set()
        wrote_any = False
        with pd.ExcelWriter(buffer, engine='openpyxl') as writer:
            for section in sections:
                table = section.get('table')
                if not isinstance(table, pd.DataFrame) or table.empty:
                    continue
                table.to_excel(
                    writer,
                    sheet_name=_safe_sheet_name(section.get('title'), used_names),
                    index=False)
                wrote_any = True
            if not wrote_any:
                pd.DataFrame({
                    '说明': ['本次选中的章节里没有可导出的统计表格。'
                             '含表格的结果（描述统计、检验结果、Table 1、回归系数表等）才会出现在这里。'],
                }).to_excel(writer, sheet_name='说明', index=False)
        return buffer.getvalue()

    @staticmethod
    def export_images_zip(sections: List[Dict], dpi: int = 300) -> bytes:
        """把选中的图打包为 ZIP：PNG + 带 dpi 元数据的 TIFF，并附一份图目录。"""
        buffer = io.BytesIO()
        entries: List[List[str]] = []
        with zipfile.ZipFile(buffer, 'w', zipfile.ZIP_DEFLATED) as archive:
            index = 0
            for section in sections:
                data = section.get('image')
                if not data:
                    continue
                index += 1
                title = str(section.get('title') or '')
                stem = _safe_file_stem(title, index)
                archive.writestr(f'图/{stem}.png', data)
                entries.append([f'{stem}.png', title, 'PNG（已按 300 dpi 生成）'])
                tiff = _to_tiff(data, dpi)
                if tiff:
                    archive.writestr(f'图/{stem}.tif', tiff)
                    entries.append([f'{stem}.tif', title,
                                    f'TIFF（含 {dpi} dpi 元数据、LZW 无损压缩）'])
            if entries:
                lines = ['文件名\t对应章节\t说明'] + ['\t'.join(row) for row in entries]
                archive.writestr('图目录.txt', '\n'.join(lines))
            else:
                archive.writestr('说明.txt', '本次选中的章节里没有图片。')
        return buffer.getvalue()

    @staticmethod
    def _add_df_table(doc, df, style=None):
        """向 docx 插入三线表；样式与数字格式按期刊模板。"""
        from modules.journal import TableStyle
        from modules.localize import translate_title

        style = style or TableStyle()
        if df is None or len(df) == 0:
            return
        df = df.reset_index(drop=True) if df.index.name else df
        n_cols = len(df.columns)
        n_rows = len(df) + 1

        table = doc.add_table(rows=n_rows, cols=n_cols)
        if style.table_center:
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = True
        ReportExporter._apply_three_line_style(table, style)

        # 表头同样按术语表翻译；变量名不在术语表里，会保持原样
        language = str(getattr(style, 'language', 'zh') or 'zh')
        columns = [translate_title(str(name), language) for name in df.columns]

        # 表头
        for j, col in enumerate(columns):
            cell = table.rows[0].cells[j]
            cell.text = col
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in paragraph.runs:
                    run.font.bold = bool(style.header_bold)
                    run.font.size = Pt(float(style.header_size_pt))

        # 数据：按列名语义格式化（p 值列 / 百分比列 / 计数列的规则不同）
        for i, (_, row) in enumerate(df.iterrows()):
            for j, val in enumerate(row):
                cell = table.rows[i + 1].cells[j]
                cell.text = format_cell(val, columns[j], style)
                for paragraph in cell.paragraphs:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for run in paragraph.runs:
                        run.font.size = Pt(float(style.body_size_pt))

    @staticmethod
    def _apply_three_line_style(table, style=None) -> None:
        """
        三线表：顶线、底线与表头下方一条细线，去掉全部竖线与内部横线。

        线宽来自期刊模板（modules/journal.py）。通过 OOXML 直接设置边框
        （python-docx 没有现成的三线表样式），并按 schema 顺序插入元素，
        避免 Word 因元素次序错误而整段忽略边框定义。
        """
        from modules.journal import TableStyle, pt_to_eighth

        style = style or TableStyle()
        top_width = pt_to_eighth(style.top_pt)
        bottom_width = pt_to_eighth(style.bottom_pt)
        header_width = pt_to_eighth(style.header_pt)

        table_properties = table._tbl.tblPr
        borders = _get_or_create_ordered(table_properties, 'w:tblBorders',
                                         _TBL_BORDERS_SUCCESSORS)
        for child in list(borders):             # 清掉旧定义，避免残留边框冲突
            borders.remove(child)
        for edge, size in (('top', top_width), ('bottom', bottom_width),
                           ('left', '0'), ('right', '0'),
                           ('insideH', '0'), ('insideV', '0')):
            element = OxmlElement(f'w:{edge}')
            if size == '0':
                element.set(qn('w:val'), 'none')
                element.set(qn('w:sz'), '0')
            else:
                element.set(qn('w:val'), 'single')
                element.set(qn('w:sz'), size)
                element.set(qn('w:color'), '000000')
            element.set(qn('w:space'), '0')
            borders.append(element)

        if not table.rows:
            return
        for cell in table.rows[0].cells:        # 表头行下方加细线
            cell_properties = cell._tc.get_or_add_tcPr()
            cell_borders = _get_or_create_ordered(cell_properties, 'w:tcBorders',
                                                  _TC_BORDERS_SUCCESSORS)
            for child in list(cell_borders):
                cell_borders.remove(child)
            bottom = OxmlElement('w:bottom')
            bottom.set(qn('w:val'), 'single')
            bottom.set(qn('w:sz'), header_width)
            bottom.set(qn('w:color'), '000000')
            bottom.set(qn('w:space'), '0')
            cell_borders.append(bottom)


class ResultCache:
    """结果缓存 - 在 session 中保存分析结果供导出使用"""

    def __init__(self):
        self.sections: List[Dict] = []

    def add_section(self, title: str, content: str = '', table=None,
                    image: bytes = None, level: int = 1):
        """
        收录一个章节。同名章节按「覆盖」处理：
        分析页在每次 rerun 都会重新生成结果，若直接追加，
        同一个「数据总览」「SPC 控制图」会在报告里重复几十遍并持续占用内存。
        """
        section = {
            'title': title,
            'content': content,
            'table': table,
            'image': image,
            'level': level
        }
        for index, existing in enumerate(self.sections):
            if existing.get('title') == title:
                self.sections[index] = section
                return
        self.sections.append(section)

    def remove(self, title: str) -> bool:
        """按标题移除一个章节，返回是否确实移除。"""
        for index, section in enumerate(self.sections):
            if section.get('title') == title:
                del self.sections[index]
                return True
        return False

    def clear(self):
        self.sections = []

    def _pick(self, sections: Optional[List[Dict]]) -> List[Dict]:
        """sections=None 表示导出全部已收录章节。"""
        return self.sections if sections is None else sections

    def export_word(self, title: str = '医疗数据分析报告',
                    sections: Optional[List[Dict]] = None,
                    table_style=None) -> bytes:
        return ReportExporter.export_word(title, self._pick(sections),
                                         table_style=table_style)

    def export_excel(self, sections: Optional[List[Dict]] = None) -> bytes:
        return ReportExporter.export_excel(self._pick(sections))

    def export_images_zip(self, sections: Optional[List[Dict]] = None,
                          dpi: int = 300) -> bytes:
        return ReportExporter.export_images_zip(self._pick(sections), dpi=dpi)
